"""Geracao reproduzivel dos graficos e dados de Backtest Analytics."""

from __future__ import annotations

from collections import Counter, OrderedDict
import math
from pathlib import Path
import shutil
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from engine.configuracao import EXPERIMENT_VERSION


COLUNAS_ROTACOES = [
    "sequence",
    "executed_at",
    "from_asset",
    "to_asset",
    "holding_days",
    "position_return",
    "realized_pnl",
    "transaction_fees",
    "sell_execution_price",
    "buy_execution_price",
    "sell_reason",
    "buy_reason",
    "subsequent_holding_days",
    "subsequent_position_return",
    "chosen_market_return",
    "counterfactual_previous_asset_return",
    "rotation_value_added",
    "rotation_regret",
    "best_alternative_asset",
    "best_alternative_return",
    "opportunity_cost",
    "maximum_favorable_excursion",
    "maximum_adverse_excursion",
    "profit_capture_ratio",
    "rotation_id",
]


def _numero(value: Any) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _texto(value: Any, default: str = "") -> str:
    if value is None:
        return default
    try:
        if pd.isna(value):
            return default
    except (TypeError, ValueError):
        pass
    text = str(value).strip()
    return text or default


def _timestamp_utc(value: Any) -> pd.Timestamp | None:
    if value is None:
        return None
    try:
        timestamp = pd.Timestamp(value)
    except (TypeError, ValueError):
        return None
    if timestamp.tzinfo is None:
        return timestamp.tz_localize("UTC")
    return timestamp.tz_convert("UTC")


def construir_rotacoes(trades: pd.DataFrame) -> pd.DataFrame:
    """Reproduz a semantica de Backtest Analytics -> Capital Rotations."""
    registros = trades.to_dict(orient="records")
    agrupados: OrderedDict[str, list[dict[str, Any]]] = OrderedDict()

    for original in registros:
        row = dict(original)
        rotation_id = _texto(row.get("rotation_id"))
        if not rotation_id:
            action = _texto(row.get("action")).upper()
            timestamp = _timestamp_utc(row.get("timestamp"))
            asset = _texto(row.get("asset"))
            if timestamp is not None and asset and action == "SELL":
                rotation_id = f"{timestamp.isoformat()}::{asset}->CASH"
                row["rotation_id"] = rotation_id
                row["rotation_from_asset"] = asset
                row["rotation_to_asset"] = "CASH"
            elif timestamp is not None and asset and action == "BUY":
                rotation_id = f"{timestamp.isoformat()}::CASH->{asset}"
                row["rotation_id"] = rotation_id
                row["rotation_from_asset"] = "CASH"
                row["rotation_to_asset"] = asset

        if rotation_id:
            agrupados.setdefault(rotation_id, []).append(row)

    output: list[dict[str, Any]] = []
    for rotation_id, rows in agrupados.items():
        sell = next(
            (row for row in rows if _texto(row.get("action")).upper() == "SELL"),
            None,
        )
        buy = next(
            (row for row in rows if _texto(row.get("action")).upper() == "BUY"),
            None,
        )
        if sell is None and buy is None:
            continue

        source = buy or sell or {}
        timestamp = _timestamp_utc(
            source.get("timestamp")
            if source.get("timestamp") is not None
            else (sell or {}).get("timestamp")
        )
        fees = [
            value
            for row in rows
            if (value := _numero(row.get("total_fee"))) is not None
        ]
        analytics_source = buy or {}

        output.append(
            {
                "rotation_id": rotation_id,
                "executed_at": timestamp.isoformat() if timestamp is not None else None,
                "from_asset": _texto(
                    source.get("rotation_from_asset")
                    or (sell or {}).get("asset"),
                    "CASH",
                ),
                "to_asset": _texto(
                    source.get("rotation_to_asset")
                    or (buy or {}).get("asset"),
                    "CASH",
                ),
                "holding_days": _numero((sell or {}).get("holding_bars")),
                "position_return": _numero((sell or {}).get("position_return")),
                "realized_pnl": _numero((sell or {}).get("realized_pnl")),
                "transaction_fees": float(sum(fees)) if fees else 0.0,
                "sell_execution_price": _numero((sell or {}).get("execution_price")),
                "buy_execution_price": _numero((buy or {}).get("execution_price")),
                "sell_reason": _texto((sell or {}).get("reason")) or None,
                "buy_reason": _texto((buy or {}).get("reason")) or None,
                "subsequent_holding_days": _numero(
                    analytics_source.get("subsequent_holding_days")
                ),
                "subsequent_position_return": _numero(
                    analytics_source.get("subsequent_position_return")
                ),
                "chosen_market_return": _numero(
                    analytics_source.get("chosen_market_return")
                ),
                "counterfactual_previous_asset_return": _numero(
                    analytics_source.get("counterfactual_previous_asset_return")
                ),
                "rotation_value_added": _numero(
                    analytics_source.get("rotation_value_added")
                ),
                "rotation_regret": _numero(
                    analytics_source.get("rotation_regret")
                ),
                "best_alternative_asset": (
                    _texto(analytics_source.get("best_alternative_asset")) or None
                ),
                "best_alternative_return": _numero(
                    analytics_source.get("best_alternative_return")
                ),
                "opportunity_cost": _numero(
                    analytics_source.get("opportunity_cost")
                ),
                "maximum_favorable_excursion": _numero(
                    analytics_source.get("maximum_favorable_excursion")
                ),
                "maximum_adverse_excursion": _numero(
                    analytics_source.get("maximum_adverse_excursion")
                ),
                "profit_capture_ratio": _numero(
                    analytics_source.get("profit_capture_ratio")
                ),
            }
        )

    output.sort(key=lambda row: row.get("executed_at") or "")
    for sequence, row in enumerate(output, start=1):
        row["sequence"] = sequence
    return pd.DataFrame(output, columns=COLUNAS_ROTACOES)


def resumir_rotacoes(rotacoes: pd.DataFrame) -> dict[str, Any]:
    if rotacoes.empty:
        return {
            "total_rotations": 0,
            "asset_to_asset_rotations": 0,
            "market_to_cash_moves": 0,
            "cash_to_market_moves": 0,
            "profitable_rotations": 0,
            "losing_rotations": 0,
            "flat_rotations": 0,
            "total_realized_pnl": 0.0,
            "total_transaction_fees": 0.0,
            "average_holding_days": None,
            "first_rotation_at": None,
            "last_rotation_at": None,
        }

    pnl = pd.to_numeric(rotacoes["realized_pnl"], errors="coerce").dropna()
    holdings = pd.to_numeric(rotacoes["holding_days"], errors="coerce").dropna()
    fees = pd.to_numeric(rotacoes["transaction_fees"], errors="coerce").fillna(0.0)
    origem = rotacoes["from_asset"].fillna("CASH").astype(str).str.upper()
    destino = rotacoes["to_asset"].fillna("CASH").astype(str).str.upper()

    return {
        "total_rotations": int(len(rotacoes)),
        "asset_to_asset_rotations": int(((origem != "CASH") & (destino != "CASH")).sum()),
        "market_to_cash_moves": int(((origem != "CASH") & (destino == "CASH")).sum()),
        "cash_to_market_moves": int(((origem == "CASH") & (destino != "CASH")).sum()),
        "profitable_rotations": int((pnl > 0).sum()),
        "losing_rotations": int((pnl < 0).sum()),
        "flat_rotations": int((pnl == 0).sum()),
        "total_realized_pnl": float(pnl.sum()) if not pnl.empty else 0.0,
        "total_transaction_fees": float(fees.sum()),
        "average_holding_days": float(holdings.mean()) if not holdings.empty else None,
        "first_rotation_at": rotacoes.iloc[0]["executed_at"],
        "last_rotation_at": rotacoes.iloc[-1]["executed_at"],
    }


def agregar_rotacoes_mensais(
    rotacoes: pd.DataFrame,
    predictions: pd.DataFrame,
) -> pd.DataFrame:
    meses: dict[str, dict[str, Any]] = {}

    def garantir_mes(timestamp: pd.Timestamp) -> dict[str, Any]:
        key = timestamp.strftime("%Y-%m")
        if key not in meses:
            meses[key] = {
                "month": key,
                "year": int(timestamp.year),
                "month_number": int(timestamp.month),
                "movement_count": 0,
                "asset_to_asset": 0,
                "market_to_cash": 0,
                "cash_to_market": 0,
                "total_realized_pnl": 0.0,
                "total_fees": 0.0,
                "profitable_exits": 0,
                "losing_exits": 0,
                "flat_exits": 0,
                "_holdings": [],
                "_bought": Counter(),
                "_sold": Counter(),
                "_equity": [],
                "session_count": 0,
                "cash_sessions": 0,
                "best_exit_asset": None,
                "best_exit_pnl": None,
                "worst_exit_asset": None,
                "worst_exit_pnl": None,
            }
        return meses[key]

    for row in rotacoes.to_dict(orient="records"):
        timestamp = _timestamp_utc(row.get("executed_at"))
        if timestamp is None:
            continue
        month = garantir_mes(timestamp)
        month["movement_count"] += 1

        source = _texto(row.get("from_asset"), "CASH").upper()
        destination = _texto(row.get("to_asset"), "CASH").upper()
        if source == "CASH" and destination != "CASH":
            month["cash_to_market"] += 1
        elif source != "CASH" and destination == "CASH":
            month["market_to_cash"] += 1
        elif source != "CASH" and destination != "CASH":
            month["asset_to_asset"] += 1

        if source != "CASH":
            month["_sold"][source] += 1
        if destination != "CASH":
            month["_bought"][destination] += 1

        holding = _numero(row.get("holding_days"))
        if holding is not None:
            month["_holdings"].append(holding)

        pnl = _numero(row.get("realized_pnl"))
        if pnl is not None:
            month["total_realized_pnl"] += pnl
            if pnl > 0:
                month["profitable_exits"] += 1
            elif pnl < 0:
                month["losing_exits"] += 1
            else:
                month["flat_exits"] += 1

            if month["best_exit_pnl"] is None or pnl > month["best_exit_pnl"]:
                month["best_exit_pnl"] = pnl
                month["best_exit_asset"] = source
            if month["worst_exit_pnl"] is None or pnl < month["worst_exit_pnl"]:
                month["worst_exit_pnl"] = pnl
                month["worst_exit_asset"] = source

        month["total_fees"] += _numero(row.get("transaction_fees")) or 0.0

    timed_rotations = []
    for row in rotacoes.to_dict(orient="records"):
        timestamp = _timestamp_utc(row.get("executed_at"))
        if timestamp is not None:
            timed_rotations.append((timestamp, row))
    timed_rotations.sort(key=lambda item: item[0])

    predictions_frame = predictions.copy()
    predictions_frame.index = pd.to_datetime(predictions_frame.index, utc=True)
    predictions_frame = predictions_frame.sort_index()

    current_asset = "CASH"
    rotation_index = 0
    for timestamp, row in predictions_frame.iterrows():
        while (
            rotation_index < len(timed_rotations)
            and timed_rotations[rotation_index][0] <= timestamp
        ):
            current_asset = _texto(
                timed_rotations[rotation_index][1].get("to_asset"),
                "CASH",
            ).upper()
            rotation_index += 1

        equity = _numero(row.get("strategy_equity"))
        if equity is None:
            continue
        month = garantir_mes(timestamp)
        month["session_count"] += 1
        if current_asset == "CASH":
            month["cash_sessions"] += 1
        month["_equity"].append(equity)

    output: list[dict[str, Any]] = []
    for key in sorted(meses):
        month = meses[key]
        holdings = month.pop("_holdings")
        bought = month.pop("_bought")
        sold = month.pop("_sold")
        equities = month.pop("_equity")

        def top_asset(counter: Counter) -> str | None:
            if not counter:
                return None
            return sorted(counter.items(), key=lambda item: (-item[1], item[0]))[0][0]

        month["average_holding_days"] = (
            float(np.mean(holdings)) if holdings else None
        )
        month["top_bought_asset"] = top_asset(bought)
        month["top_sold_asset"] = top_asset(sold)
        month["first_equity"] = equities[0] if equities else None
        month["last_equity"] = equities[-1] if equities else None
        month["equity_return"] = (
            month["last_equity"] / month["first_equity"] - 1.0
            if month["first_equity"] not in {None, 0}
            and month["last_equity"] is not None
            else None
        )
        month["market_exposure"] = (
            (month["session_count"] - month["cash_sessions"])
            / month["session_count"]
            if month["session_count"]
            else None
        )
        output.append(month)

    return pd.DataFrame(output)


def construir_matriz_transicoes(rotacoes: pd.DataFrame) -> pd.DataFrame:
    if rotacoes.empty:
        return pd.DataFrame(
            columns=[
                "from_asset",
                "to_asset",
                "rotations",
                "profitable_rotations",
                "win_rate",
                "total_realized_pnl",
                "average_position_return",
                "transaction_fees",
            ]
        )

    rows: list[dict[str, Any]] = []
    for (source, destination), group in rotacoes.groupby(
        ["from_asset", "to_asset"],
        dropna=False,
    ):
        pnl = pd.to_numeric(group["realized_pnl"], errors="coerce").dropna()
        returns = pd.to_numeric(group["position_return"], errors="coerce").dropna()
        fees = pd.to_numeric(group["transaction_fees"], errors="coerce").fillna(0.0)
        rows.append(
            {
                "from_asset": source,
                "to_asset": destination,
                "rotations": int(len(group)),
                "profitable_rotations": int((pnl > 0).sum()),
                "win_rate": (
                    float((pnl > 0).sum() / len(pnl))
                    if not pnl.empty
                    else None
                ),
                "total_realized_pnl": float(pnl.sum()) if not pnl.empty else 0.0,
                "average_position_return": (
                    float(returns.mean()) if not returns.empty else None
                ),
                "transaction_fees": float(fees.sum()),
            }
        )
    return pd.DataFrame(rows).sort_values(
        ["rotations", "from_asset", "to_asset"],
        ascending=[False, True, True],
        ignore_index=True,
    )


def calcular_pnl_realizado_mensal(trades: pd.DataFrame) -> pd.DataFrame:
    rows = trades.copy()
    rows["timestamp"] = pd.to_datetime(rows["timestamp"], utc=True)
    rows["action"] = rows["action"].astype(str).str.upper()
    rows = rows[rows["action"].isin(["SELL", "FINAL_SELL"])].copy()
    rows["realized_pnl"] = pd.to_numeric(rows["realized_pnl"], errors="coerce")
    rows["month"] = rows["timestamp"].dt.strftime("%Y-%m")
    return (
        rows.groupby("month", as_index=False)["realized_pnl"]
        .sum()
        .sort_values("month", ignore_index=True)
    )


def calcular_retornos_mensais(predictions: pd.DataFrame) -> pd.DataFrame:
    frame = predictions.copy()
    frame.index = pd.to_datetime(frame.index, utc=True)
    frame = frame.sort_index()
    frame["month"] = frame.index.strftime("%Y-%m")
    monthly = frame.groupby("month", sort=True).tail(1).copy()
    monthly = monthly.sort_index()

    output = pd.DataFrame(
        {
            "month": monthly["month"].to_numpy(),
            "simulation_return": pd.to_numeric(
                monthly["strategy_equity"],
                errors="coerce",
            ).pct_change().to_numpy(),
            "reference_return": pd.to_numeric(
                monthly["buy_hold_equity"],
                errors="coerce",
            ).pct_change().to_numpy(),
        }
    )
    output["excess_return"] = (
        output["simulation_return"] - output["reference_return"]
    )
    return output.iloc[1:].reset_index(drop=True)


def _matriz_ano_mes(
    data: pd.DataFrame,
    *,
    value_column: str,
) -> pd.DataFrame:
    if data.empty:
        columns = [*range(1, 13), "Total"]
        return pd.DataFrame(columns=columns)

    frame = data.copy()
    frame["year"] = frame["month"].str.slice(0, 4).astype(int)
    frame["month_number"] = frame["month"].str.slice(5, 7).astype(int)
    matrix = frame.pivot_table(
        index="year",
        columns="month_number",
        values=value_column,
        aggfunc="sum",
        fill_value=0.0,
    )
    for month in range(1, 13):
        if month not in matrix.columns:
            matrix[month] = 0.0
    matrix = matrix[[*range(1, 13)]].sort_index()
    matrix["Total"] = matrix.sum(axis=1)
    total = matrix.sum(axis=0)
    total.name = "Total"
    matrix = pd.concat([matrix, total.to_frame().T])
    return matrix


def _valor_compacto(value: float) -> str:
    absolute = abs(float(value))
    if absolute >= 1_000_000:
        return f"{value / 1_000_000:.1f}M"
    if absolute >= 1_000:
        return f"{value / 1_000:.1f}K"
    return f"{value:.0f}"


def _salvar_heatmap(
    matrix: pd.DataFrame,
    *,
    title: str,
    png_path: Path,
    svg_path: Path,
    kind: str,
) -> None:
    if matrix.empty:
        return

    values = matrix.to_numpy(dtype=float)
    display = values.copy()
    color_values = values.copy()

    if color_values.shape[0] > 1 and color_values.shape[1] > 1:
        core = color_values[:-1, :-1]
        core_scale = float(np.nanmax(np.abs(core))) if core.size else 0.0
        if core_scale <= 0:
            core_scale = 1.0
        color_values[-1, :] = np.sign(color_values[-1, :]) * core_scale * 0.55
        color_values[:, -1] = np.sign(color_values[:, -1]) * core_scale * 0.55

    labels = [
        "Jan",
        "Fev",
        "Mar",
        "Abr",
        "Mai",
        "Jun",
        "Jul",
        "Ago",
        "Set",
        "Out",
        "Nov",
        "Dez",
        "Total",
    ]
    ylabels = [str(item) for item in matrix.index]

    fig = plt.figure(figsize=(15, 6.8))
    ax = fig.add_axes([0.07, 0.16, 0.88, 0.72])

    if kind in {"return", "pnl"}:
        limit = float(np.nanmax(np.abs(color_values)))
        if limit <= 0:
            limit = 1.0
        image = ax.imshow(
            color_values,
            aspect="auto",
            vmin=-limit,
            vmax=limit,
            cmap="RdYlGn",
        )
    else:
        image = ax.imshow(color_values, aspect="auto", cmap="Blues")

    ax.set_xticks(range(len(labels)), labels)
    ax.set_yticks(range(len(ylabels)), ylabels)
    ax.set_title(title)
    ax.set_xlabel("Mes")
    ax.set_ylabel("Ano")

    for row in range(display.shape[0]):
        for col in range(display.shape[1]):
            value = display[row, col]
            if kind == "return":
                text = f"{value:+.1%}"
            elif kind == "pnl":
                text = _valor_compacto(value)
            else:
                text = str(int(round(value)))
            ax.text(col, row, text, ha="center", va="center", fontsize=8)

    fig.colorbar(image, ax=ax)
    fig.savefig(png_path, dpi=180, bbox_inches="tight")
    fig.savefig(svg_path, bbox_inches="tight")
    plt.close(fig)


def _formatar_excel(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="E5E7EB")
    header_font = Font(bold=True)

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for column_cells in sheet.columns:
            letter = column_cells[0].column_letter
            width = 10
            for cell in column_cells[:200]:
                if cell.value is None:
                    continue
                width = max(width, min(38, len(str(cell.value)) + 2))
            sheet.column_dimensions[letter].width = width

        headers = {
            cell.column: str(cell.value or "").lower()
            for cell in sheet[1]
        }
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                header = headers.get(cell.column, "")
                if any(
                    token in header
                    for token in (
                        "return",
                        "rate",
                        "exposure",
                        "regret",
                        "opportunity_cost",
                        "excursion",
                        "capture_ratio",
                    )
                ):
                    cell.number_format = "0.00%"
                elif any(
                    token in header
                    for token in (
                        "pnl",
                        "fee",
                        "equity",
                        "price",
                        "capital",
                    )
                ):
                    cell.number_format = "$#,##0.00"

    workbook.save(path)


def _salvar_excel(
    path: Path,
    *,
    manifest: dict[str, Any],
    control: dict[str, pd.DataFrame],
    soft: dict[str, pd.DataFrame],
    control_summary: dict[str, Any],
    soft_summary: dict[str, Any],
) -> None:
    resumo = pd.DataFrame(
        [
            {
                "metric": key,
                "control": control_summary.get(key),
                "soft_horizon_consensus": soft_summary.get(key),
            }
            for key in control_summary
        ]
    )
    metadados = pd.DataFrame(
        [
            {"field": "experiment_version", "value": EXPERIMENT_VERSION},
            {
                "field": "snapshot_sha256",
                "value": manifest.get("snapshot_sha256"),
            },
            {"field": "database_access", "value": False},
            {"field": "source", "value": "Backtest Analytics local"},
        ]
    )

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        metadados.to_excel(writer, sheet_name="Metadados", index=False)
        resumo.to_excel(writer, sheet_name="Resumo", index=False)
        control["rotacoes"].to_excel(
            writer,
            sheet_name="Control Rotacoes",
            index=False,
        )
        soft["rotacoes"].to_excel(
            writer,
            sheet_name="Soft Rotacoes",
            index=False,
        )
        control["rotacoes_mensais"].to_excel(
            writer,
            sheet_name="Control Rotacoes Mensais",
            index=False,
        )
        soft["rotacoes_mensais"].to_excel(
            writer,
            sheet_name="Soft Rotacoes Mensais",
            index=False,
        )
        control["transicoes"].to_excel(
            writer,
            sheet_name="Control Transicoes",
            index=False,
        )
        soft["transicoes"].to_excel(
            writer,
            sheet_name="Soft Transicoes",
            index=False,
        )
        control["retornos_mensais"].to_excel(
            writer,
            sheet_name="Control Retornos",
            index=False,
        )
        soft["retornos_mensais"].to_excel(
            writer,
            sheet_name="Soft Retornos",
            index=False,
        )
        control["pnl_mensal"].to_excel(
            writer,
            sheet_name="Control PnL Mensal",
            index=False,
        )
        soft["pnl_mensal"].to_excel(
            writer,
            sheet_name="Soft PnL Mensal",
            index=False,
        )
    _formatar_excel(path)


def _gerar_variante(
    output_dir: Path,
    *,
    slug: str,
    label: str,
    result: Any,
) -> tuple[dict[str, pd.DataFrame], dict[str, Path], dict[str, Any]]:
    rotacoes = construir_rotacoes(result.trades)
    rotacoes_mensais = agregar_rotacoes_mensais(
        rotacoes,
        result.predictions,
    )
    transicoes = construir_matriz_transicoes(rotacoes)
    pnl_mensal = calcular_pnl_realizado_mensal(result.trades)
    retornos_mensais = calcular_retornos_mensais(result.predictions)
    resumo = resumir_rotacoes(rotacoes)

    frames = {
        "rotacoes": rotacoes,
        "rotacoes_mensais": rotacoes_mensais,
        "transicoes": transicoes,
        "pnl_mensal": pnl_mensal,
        "retornos_mensais": retornos_mensais,
    }

    paths = {
        "rotacoes": output_dir / f"capital_rotations_{slug}.csv",
        "rotacoes_mensais": (
            output_dir / f"capital_rotations_monthly_{slug}.csv"
        ),
        "transicoes": (
            output_dir / f"capital_rotations_transition_matrix_{slug}.csv"
        ),
        "pnl_mensal": output_dir / f"monthly_realized_pnl_{slug}.csv",
        "retornos_mensais": output_dir / f"monthly_returns_{slug}.csv",
    }

    rotacoes.to_csv(paths["rotacoes"], index=False)
    rotacoes_mensais.to_csv(paths["rotacoes_mensais"], index=False)
    transicoes.to_csv(paths["transicoes"], index=False)
    pnl_mensal.to_csv(paths["pnl_mensal"], index=False)
    retornos_mensais.to_csv(paths["retornos_mensais"], index=False)

    pnl_matrix = _matriz_ano_mes(
        pnl_mensal,
        value_column="realized_pnl",
    )
    _salvar_heatmap(
        pnl_matrix,
        title=f"Backtest Analytics · P/L realizado mensal · {label}",
        png_path=output_dir / f"monthly_realized_pnl_heatmap_{slug}.png",
        svg_path=output_dir / f"monthly_realized_pnl_heatmap_{slug}.svg",
        kind="pnl",
    )

    rotation_matrix = _matriz_ano_mes(
        rotacoes_mensais.rename(
            columns={"movement_count": "value"}
        )[["month", "value"]],
        value_column="value",
    )
    rotation_matrix.to_csv(
        output_dir / f"capital_rotations_heatmap_{slug}.csv",
        index=True,
        index_label="year",
    )
    _salvar_heatmap(
        rotation_matrix,
        title=f"Backtest Analytics · Capital Rotations · {label}",
        png_path=output_dir / f"capital_rotations_heatmap_{slug}.png",
        svg_path=output_dir / f"capital_rotations_heatmap_{slug}.svg",
        kind="count",
    )

    for metric, suffix in (
        ("simulation_return", "simulation"),
        ("reference_return", "reference"),
        ("excess_return", "excess"),
    ):
        matrix = _matriz_ano_mes(
            retornos_mensais,
            value_column=metric,
        )
        matrix.to_csv(
            output_dir
            / f"monthly_return_heatmap_{slug}_{suffix}.csv",
            index=True,
            index_label="year",
        )
        _salvar_heatmap(
            matrix,
            title=(
                "Backtest Analytics · Retorno mensal "
                f"({suffix}) · {label}\n"
                "Total = soma aritmetica dos retornos mensais"
            ),
            png_path=(
                output_dir
                / f"monthly_return_heatmap_{slug}_{suffix}.png"
            ),
            svg_path=(
                output_dir
                / f"monthly_return_heatmap_{slug}_{suffix}.svg"
            ),
            kind="return",
        )

    return frames, paths, resumo


def gerar_analises_backtest(
    output_dir: Path,
    *,
    manifest: dict[str, Any],
    control_result: Any,
    soft_result: Any,
) -> dict[str, Path]:
    """Gera todos os graficos e dados analiticos da reproducao.

    A pasta de graficos e recriada integralmente para impedir mistura entre
    artefatos de execucoes diferentes.
    """
    graficos_dir = output_dir / "graficos"
    if graficos_dir.exists():
        shutil.rmtree(graficos_dir)
    graficos_dir.mkdir(parents=True, exist_ok=True)

    control, control_paths, control_summary = _gerar_variante(
        graficos_dir,
        slug="control",
        label="Control",
        result=control_result,
    )
    soft, soft_paths, soft_summary = _gerar_variante(
        graficos_dir,
        slug="soft",
        label="Soft Horizon Consensus",
        result=soft_result,
    )

    excel_path = graficos_dir / "backtest_analytics.xlsx"
    _salvar_excel(
        excel_path,
        manifest=manifest,
        control=control,
        soft=soft,
        control_summary=control_summary,
        soft_summary=soft_summary,
    )

    print(
        "[analytics] "
        f"control_rotations={control_summary['total_rotations']} "
        f"soft_rotations={soft_summary['total_rotations']}",
        flush=True,
    )
    print(f"[analytics] dir={graficos_dir}", flush=True)

    return {
        **{f"control_{key}": value for key, value in control_paths.items()},
        **{f"soft_{key}": value for key, value in soft_paths.items()},
        "backtest_analytics_xlsx": excel_path,
        "graficos_dir": graficos_dir,
    }
