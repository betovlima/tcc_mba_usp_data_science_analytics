from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "output"
DEFAULT_DEST = ROOT / "analysis" / "tcc_backtest_output_analysis.xlsx"


def args():
    parser = argparse.ArgumentParser(description="Build the Excel audit workbook from output/.")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--destination", type=Path, default=DEFAULT_DEST)
    return parser.parse_args()


def load_data(output_dir: Path):
    required = [
        "backtest_result.json",
        "equity_curve.csv",
        "folds.csv",
        "trades.csv",
        "summary.txt",
        "market_data.csv",
    ]
    missing = [name for name in required if not (output_dir / name).exists()]
    if missing:
        raise FileNotFoundError("Missing files: " + ", ".join(missing))

    result = json.loads((output_dir / "backtest_result.json").read_text(encoding="utf-8"))
    equity = pd.read_csv(output_dir / "equity_curve.csv")
    folds = pd.read_csv(output_dir / "folds.csv")
    trades = pd.read_csv(output_dir / "trades.csv")
    market = pd.read_csv(output_dir / "market_data.csv")
    summary = (output_dir / "summary.txt").read_text(encoding="utf-8")

    market["timestamp"] = pd.to_datetime(market["timestamp"], utc=True)
    return result, equity, folds, trades, market, summary


def build_cycles(trades: pd.DataFrame) -> pd.DataFrame:
    buy_columns = [
        "asset",
        "timestamp",
        "best_alternative_asset",
        "best_alternative_return",
        "opportunity_cost",
    ]
    sell_columns = [
        "entry_timestamp",
        "timestamp",
        "asset",
        "entry_price",
        "execution_price",
        "quantity",
        "realized_pnl",
        "position_return",
        "holding_bars",
        "total_fee",
        "walk_forward_fold",
        "rotation_id",
        "position_entry_score",
        "current_score",
        "best_score",
        "best_vs_current_gap",
        "effective_switch_margin",
        "current_asset_rank",
        "maximum_favorable_excursion",
        "maximum_adverse_excursion",
        "profit_capture_ratio",
    ]

    buy_columns = [column for column in buy_columns if column in trades.columns]
    sell_columns = [column for column in sell_columns if column in trades.columns]

    buys = trades.loc[trades["action"].eq("BUY"), buy_columns].copy()
    if "timestamp" in buys.columns:
        buys = buys.rename(columns={"timestamp": "entry_timestamp"})

    sells = trades.loc[trades["action"].isin(["SELL", "FINAL_SELL"]), sell_columns].copy()
    cycles = sells.merge(buys, on=["asset", "entry_timestamp"], how="left")
    cycles = cycles.rename(
        columns={
            "timestamp": "exit_timestamp",
            "execution_price": "exit_price",
            "total_fee": "exit_fee",
            "walk_forward_fold": "fold",
            "position_entry_score": "entry_score",
            "maximum_favorable_excursion": "mfe",
            "maximum_adverse_excursion": "mae",
        }
    )
    if "position_return" in cycles.columns:
        cycles["result"] = cycles["position_return"].map(
            lambda value: "WIN" if value > 0 else "LOSS" if value < 0 else "FLAT"
        )
    return cycles


def build_monthly(equity: pd.DataFrame) -> pd.DataFrame:
    frame = equity.copy()
    timestamp_column = frame.columns[0]
    frame[timestamp_column] = (
        pd.to_datetime(frame[timestamp_column], utc=True).dt.tz_localize(None)
    )
    frame["month"] = frame[timestamp_column].dt.to_period("M").astype(str)

    monthly = frame.groupby("month", as_index=False).agg(
        strategy_month_end=("strategy_equity", "last"),
        benchmark_month_end=("buy_hold_equity", "last"),
    )
    monthly["strategy_monthly_return"] = monthly["strategy_month_end"].pct_change()
    monthly["benchmark_monthly_return"] = monthly["benchmark_month_end"].pct_change()
    monthly["excess_monthly_return"] = (
        monthly["strategy_monthly_return"] - monthly["benchmark_monthly_return"]
    )
    return monthly


def build_asset_summary(market: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for symbol, frame in market.groupby("symbol", sort=False):
        frame = frame.sort_values("timestamp").copy()
        timestamps = pd.to_datetime(frame["timestamp"], utc=True).dt.tz_localize(None)
        first_close = float(frame.iloc[0]["close"])
        last_close = float(frame.iloc[-1]["close"])
        rows.append(
            {
                "asset": symbol,
                "rows": len(frame),
                "start": timestamps.iloc[0],
                "end": timestamps.iloc[-1],
                "first_close": first_close,
                "last_close": last_close,
                "close_return": last_close / first_close - 1.0,
            }
        )
    return pd.DataFrame(rows)


def flatten(value, prefix=""):
    rows = []
    if isinstance(value, dict):
        for key, child in value.items():
            rows.extend(flatten(child, f"{prefix}.{key}" if prefix else str(key)))
    elif isinstance(value, list):
        rows.append((prefix, "list", json.dumps(value, ensure_ascii=False, default=str)))
    else:
        rows.append((prefix, type(value).__name__, value))
    return rows


def style_asset_sheet(ws):
    last_row = ws.max_row
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

    ws["G1"] = "Daily Return"
    ws["H1"] = "Running Peak"
    ws["I1"] = "Drawdown"

    for row in range(2, last_row + 1):
        if row == 2:
            ws[f"H{row}"] = f"=E{row}"
            ws[f"I{row}"] = 0
        else:
            ws[f"G{row}"] = f"=E{row}/E{row-1}-1"
            ws[f"H{row}"] = f"=MAX(H{row-1},E{row})"
            ws[f"I{row}"] = f"=E{row}/H{row}-1"

    for row in range(2, last_row + 1):
        ws[f"A{row}"].number_format = "yyyy-mm-dd"
        for column in "BCDEH":
            ws[f"{column}{row}"].number_format = "0.0000"
        ws[f"F{row}"].number_format = "#,##0"
        ws[f"G{row}"].number_format = "0.0000%"
        ws[f"I{row}"].number_format = "0.0000%"

    if last_row >= 2:
        ws.conditional_formatting.add(
            f"I2:I{last_row}",
            ColorScaleRule(
                start_type="min",
                start_color="F8696B",
                mid_type="percentile",
                mid_value=50,
                mid_color="FFEB84",
                end_type="max",
                end_color="63BE7B",
            ),
        )


def style_workbook(path: Path, result: dict, summary: str, asset_names: list[str]):
    wb = load_workbook(path)
    blue, navy, white = "2F75B5", "17365D", "FFFFFF"

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor=blue)
            cell.font = Font(color=white, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col in range(1, min(ws.max_column, 20) + 1):
            width = max(
                11,
                min(
                    28,
                    max(
                        len(str(ws.cell(row, col).value or ""))
                        for row in range(1, min(ws.max_row, 40) + 1)
                    )
                    + 2,
                ),
            )
            ws.column_dimensions[get_column_letter(col)].width = width

    for symbol in asset_names:
        if symbol in wb.sheetnames:
            style_asset_sheet(wb[symbol])

    guide = wb["00_Guia"]
    guide.insert_rows(1, 1)
    guide.merge_cells("A1:F1")
    guide["A1"] = "Auditoria do backtest no Excel"
    guide["A1"].fill = PatternFill("solid", fgColor=navy)
    guide["A1"].font = Font(color=white, bold=True, size=15)

    equity_ws = wb["02_Equity"]
    last_eq = equity_ws.max_row
    base_headers = equity_ws.max_column
    derived = [
        "Strategy Daily Return",
        "Benchmark Daily Return",
        "Strategy Peak",
        "Strategy Drawdown",
    ]
    for offset, name in enumerate(derived, start=1):
        equity_ws.cell(1, base_headers + offset, name)
    for row in range(2, last_eq + 1):
        first = base_headers + 1
        if row > 2:
            equity_ws.cell(row, first, f"=B{row}/B{row-1}-1")
            equity_ws.cell(row, first + 1, f"=C{row}/C{row-1}-1")
        equity_ws.cell(
            row,
            first + 2,
            f"=MAX(10000,B{row})"
            if row == 2
            else f"=MAX({get_column_letter(first+2)}{row-1},B{row})",
        )
        equity_ws.cell(row, first + 3, f"=B{row}/{get_column_letter(first+2)}{row}-1")

    kpi = wb["01_KPIs"]
    kpi["A1"] = "Métrica"
    kpi["B1"] = "Motor / JSON"
    kpi["C1"] = "Excel"
    kpi["D1"] = "Diferença"
    kpi["E1"] = "Status"
    kpi["F1"] = "Reconstrução"

    metrics = result["metrics"]
    last_trades = wb["04_Trades"].max_row
    last_cycles = wb["05_Cycles"].max_row
    fold_rows = list(metrics.get("walk_forward_folds") or [])
    worst_fold = min(float(item["strategy_return"]) for item in fold_rows) if fold_rows else None

    rows = [
        ("Initial Capital", metrics.get("initial_capital"), "='03_Folds'!K2", "Capital inicial."),
        ("Final Capital", metrics.get("strategy_ending_capital"), f"='02_Equity'!B{last_eq}", "Última equity."),
        ("Strategy Return", metrics.get("strategy_return"), "=C3/C2-1", "Final/inicial - 1."),
        ("CAGR", metrics.get("strategy_cagr"), None, "Motor."),
        ("Sharpe", metrics.get("strategy_sharpe"), None, "Motor."),
        ("Max Drawdown", metrics.get("strategy_maximum_drawdown"), None, "Motor."),
        ("Buys", metrics.get("simulated_buys"), f'=COUNTIF(\'04_Trades\'!B2:B{last_trades},"BUY")', "Compras."),
        ("Sells", metrics.get("simulated_sells"), f'=COUNTIF(\'04_Trades\'!B2:B{last_trades},"SELL")+COUNTIF(\'04_Trades\'!B2:B{last_trades},"FINAL_SELL")', "Vendas."),
        ("Rotations", metrics.get("capital_rotations"), f'=COUNTIF(\'04_Trades\'!B2:B{last_trades},"SELL")', "Rotações."),
        ("Avg Holding Bars", metrics.get("average_holding_bars"), f"=AVERAGE('05_Cycles'!I2:I{last_cycles})", "Holding médio."),
        ("Total Fees", metrics.get("total_transaction_fees"), None, "Custos do motor."),
        ("Worst Fold Return", worst_fold, None, "Pior fold."),
    ]

    for row_index, (name, engine, formula, note) in enumerate(rows, start=2):
        kpi.cell(row_index, 1, name)
        kpi.cell(row_index, 2, engine)
        if formula:
            kpi.cell(row_index, 3, formula)
            kpi.cell(row_index, 4, f"=C{row_index}-B{row_index}")
            kpi.cell(row_index, 5, f'=IF(ABS(D{row_index})<1E-8,"OK","CHECK")')
        else:
            kpi.cell(row_index, 3, engine)
            kpi.cell(row_index, 4, 0)
            kpi.cell(row_index, 5, "REFERENCE")
        kpi.cell(row_index, 6, note)

    chart = LineChart()
    chart.title = "Capital: Strategy vs Benchmark"
    chart.add_data(
        Reference(equity_ws, min_col=2, max_col=3, min_row=1, max_row=last_eq),
        titles_from_data=True,
    )
    chart.height, chart.width = 8, 16
    kpi.add_chart(chart, "H2")

    raw = wb["09_JSON"]
    raw.cell(raw.max_row + 2, 1, "summary.txt")
    raw.cell(raw.max_row + 1, 1, summary)

    if hasattr(wb, "calculation"):
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"

    wb.save(path)


def main():
    config = args()
    result, equity, folds, trades, market, summary = load_data(config.output_dir)
    cycles = build_cycles(trades)
    monthly = build_monthly(equity)
    asset_summary = build_asset_summary(market)

    guide = pd.DataFrame(
        [
            ["market_data.csv", "Snapshot Yahoo", "OHLCV usado pelo motor", "Sim"],
            ["backtest_result.json", "Resultado canônico", "Métricas e metadados", "Parcial"],
            ["equity_curve.csv", "Curva de capital", "CAGR, Sharpe, drawdown", "Sim"],
            ["folds.csv", "Walk-forward", "Retorno e compound por fold", "Sim"],
            ["trades.csv", "Livro-razão", "PnL, custos e diagnósticos", "Sim após decisão"],
            ["summary.txt", "Resumo", "Conferência", "Não necessário"],
        ],
        columns=["Arquivo", "Papel", "Auditoria", "Excel"],
    )

    dictionary = pd.DataFrame(
        [
            ["market_data.csv", "Séries temporais OHLCV baixadas do Yahoo e usadas no backtest"],
            ["strategy_equity", "Capital ao fim da sessão"],
            ["effective_switch_margin", "Margem mínima para rotação"],
            ["current_score", "Utility da posição atual"],
            ["best_score", "Utility do melhor candidato"],
            ["maximum_favorable_excursion", "MFE: maior excursão favorável"],
            ["maximum_adverse_excursion", "MAE: maior excursão adversa"],
            ["profit_capture_ratio", "Fração do MFE capturada"],
        ],
        columns=["Campo", "Significado"],
    )

    flat = pd.DataFrame(flatten(result), columns=["JSON Path", "Type", "Value"])

    config.destination.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(config.destination, engine="openpyxl") as writer:
        guide.to_excel(writer, sheet_name="00_Guia", index=False)
        pd.DataFrame(
            columns=["Métrica", "Motor / JSON", "Excel", "Diferença", "Status", "Reconstrução"]
        ).to_excel(writer, sheet_name="01_KPIs", index=False)
        equity.to_excel(writer, sheet_name="02_Equity", index=False)
        folds.to_excel(writer, sheet_name="03_Folds", index=False)
        trades.to_excel(writer, sheet_name="04_Trades", index=False)
        cycles.to_excel(writer, sheet_name="05_Cycles", index=False)
        asset_summary.to_excel(writer, sheet_name="06_Assets", index=False)
        monthly.to_excel(writer, sheet_name="07_Monthly", index=False)
        dictionary.to_excel(writer, sheet_name="08_Dictionary", index=False)
        flat.to_excel(writer, sheet_name="09_JSON", index=False)
        pd.DataFrame(columns=["Item", "Valor"]).to_excel(
            writer, sheet_name="10_Reconciliation", index=False
        )
        asset_summary.to_excel(writer, sheet_name="11_Ativos", index=False)

        asset_order = [str(symbol) for symbol in result.get("assets", [])]
        if not asset_order:
            asset_order = list(dict.fromkeys(market["symbol"].astype(str)))

        for symbol in asset_order:
            series = market.loc[market["symbol"].eq(symbol)].copy()
            series = series.sort_values("timestamp")
            series = series[["timestamp", "open", "high", "low", "close", "volume"]]
            series.columns = ["Date", "Open", "High", "Low", "Close", "Volume"]
            series["Date"] = pd.to_datetime(series["Date"], utc=True).dt.tz_localize(None)
            series.to_excel(writer, sheet_name=symbol[:31], index=False)

    asset_names = [str(symbol)[:31] for symbol in result.get("assets", [])]
    style_workbook(config.destination, result, summary, asset_names)
    print(f"Excel generated: {config.destination}")


if __name__ == "__main__":
    main()
